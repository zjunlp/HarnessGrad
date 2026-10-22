"""Authored formula-only control. No dataset records or released workbooks are read."""
import errno
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import uuid
import zipfile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
import openpyxl
from data.scaffolds.scorer_input import inventory_digest


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    work = Path(sys.argv[1])
    work.mkdir(parents=True)
    print('label=synthetic offline control; not dataset measurement or official scorer execution', flush=True)
    print('work=' + str(work), flush=True)
    print('netns=' + os.readlink('/proc/self/ns/net'), flush=True)
    print('parent_netns=' + os.environ['SPREADRT_PARENT_NETNS'], flush=True)
    assert os.readlink('/proc/self/ns/net') != os.environ['SPREADRT_PARENT_NETNS']
    print('route_table=' + repr(Path('/proc/net/route').read_text()), flush=True)
    for family, host in ((socket.AF_INET, '192.0.2.1'), (socket.AF_INET6, '2001:db8::1')):
        with socket.socket(family, socket.SOCK_DGRAM) as s:
            try:
                s.connect((host, 9))
            except OSError as exc:
                print(f'network_probe family={family.name} errno={exc.errno} message={exc.strerror}', flush=True)
                assert exc.errno in (errno.ENETUNREACH, errno.EPERM, errno.EACCES)
            else:
                raise AssertionError('external route unexpectedly available')
    version = subprocess.run(['/usr/bin/soffice', '--version'], capture_output=True, text=True, check=True)
    print('engine_version_stdout=' + repr(version.stdout), flush=True)
    print('engine_version_stderr=' + repr(version.stderr), flush=True)
    wb = openpyxl.Workbook()
    wb.active.title = 'Synthetic'
    wb.active['A1'] = 2
    wb.active['A2'] = 3
    wb.active['B1'] = '=SUM(A1:A2)'
    wb.active['B2'] = '=B1*2'
    candidate = work / 'result.xlsx'
    wb.save(candidate)
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(candidate) as z:
        entries = {name: z.read(name) for name in z.namelist()}
    xml = ET.fromstring(entries['xl/worksheets/sheet1.xml'])
    for name, stale in (('B1', 999), ('B2', 888)):
        cell = xml.find(f"s:sheetData/s:row/s:c[@r='{name}']", ns)
        cell.find('s:v', ns).text = str(stale)
    entries['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    with zipfile.ZipFile(candidate, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, content in entries.items():
            z.writestr(name, content)
    golden = work / 'golden.xlsx'
    wb.active['B1'] = 5
    wb.active['B2'] = 10
    wb.save(golden)
    wb.close()
    task = {'task_id': 'spreadrt-authored-formulas', 'answer_position': 'Synthetic!B1:B2',
            'instruction_type': 'synthetic'}
    (work / 'task.json').write_text(json.dumps(task))
    (work / 'answer.txt').write_text('result.xlsx\n')
    check_path = ROOT / 'data/datasets/spreadsheetbench/verify/check.py'
    spec = importlib.util.spec_from_file_location('spreadrt_check', check_path)
    check = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(check)
    cached = openpyxl.load_workbook(candidate, data_only=True)
    before = [cached.active['B1'].value, cached.active['B2'].value]
    cached.close()
    print('authored_formulas=["=SUM(A1:A2)", "=B1*2"]', flush=True)
    print('stale_cached_values=' + json.dumps(before), flush=True)
    initial = check.compare_workbooks(str(golden), str(candidate), 'synthetic', task['answer_position'])
    print('before_compare=' + repr(initial), flush=True)
    assert before == [999, 888] and initial[0] is False
    files = {name: sha(work / name) for name in ('golden.xlsx', 'task.json')}
    argv = [sys.executable, str(check_path), '--workdir', str(work),
            '--reward-file', str(work / 'reward.txt'), '--evidence-file', str(work / 'evidence.json'),
            '--task-id', task['task_id'], '--scorer-input-sha256', inventory_digest(files),
            '--scorer-input-files', json.dumps(files, sort_keys=True)]
    print('scorer_argv=' + json.dumps(argv), flush=True)
    proc = subprocess.run(argv, capture_output=True, text=True, cwd=ROOT)
    print('scorer_exit=' + str(proc.returncode), flush=True)
    print('scorer_stdout=' + repr(proc.stdout), flush=True)
    print('scorer_stderr=' + repr(proc.stderr), flush=True)
    evidence = json.loads((work / 'evidence.json').read_text())
    print('scorer_evidence=' + json.dumps(evidence, sort_keys=True), flush=True)
    print('reward_file=' + (repr((work / 'reward.txt').read_text())
                           if (work / 'reward.txt').exists() else 'ABSENT'), flush=True)
    cached = openpyxl.load_workbook(candidate, data_only=True)
    after = [cached.active['B1'].value, cached.active['B2'].value]
    cached.close()
    print('recalculated_cached_values=' + json.dumps(after), flush=True)
    assert proc.returncode == 0 and after == [5, 10]
    assert evidence['measured'] is True and evidence['recalculated'] is True
    assert evidence['scorer_input']['verified'] is True
    assert evidence['score'] == 1.0 and (work / 'reward.txt').read_text() == '1\n'
    print('synthetic_control=PASSED', flush=True)


if __name__ == '__main__':
    main()
